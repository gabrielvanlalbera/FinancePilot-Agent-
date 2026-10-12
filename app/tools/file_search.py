
from pathlib import Path
from openpyxl import load_workbook
import re

PROJECT_ROOT = Path(__file__).resolve().parents[2]
DATA_DIR = PROJECT_ROOT / "data"

BANK_FILE = DATA_DIR / "bank_transactions" / "transactions.xlsx"
INVOICE_FILE = DATA_DIR / "invoice" / "invoice.xlsx"
POLICY_FILE = DATA_DIR / "policies" / "finance_policy.txt"


def normalize(value):
    """Normalize text for comparisons."""
    return re.sub(r"\s+", " ", str(value or "")).strip().casefold()


def read_excel_records(file_path):
    """Read an Excel worksheet into a list of dictionaries."""
    if not file_path.exists():
        return []

    workbook = load_workbook(file_path, read_only=True, data_only=True)

    try:
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        headers = next(rows, None)

        if not headers:
            return []

        headers = [
            str(header).strip() if header is not None else ""
            for header in headers
        ]

        records = []

        for row in rows:
            record = {}

            for header, value in zip(headers, row):
                if value is not None and hasattr(value, "isoformat"):
                    value = (
                        value.isoformat(sep=" ")
                        if hasattr(value, "hour")
                        else value.isoformat()
                    )

                record[header] = value

            if any(value is not None for value in record.values()):
                records.append(record)

        return records
    finally:
        workbook.close()


def load_datasets():
    """Load the available bank, invoice, and policy data."""
    datasets = {}

    if BANK_FILE.exists():
        datasets["Bank Transactions"] = {
            "file": BANK_FILE.name,
            "type": "table",
            "data": read_excel_records(BANK_FILE),
        }

    if INVOICE_FILE.exists():
        datasets["Invoices"] = {
            "file": INVOICE_FILE.name,
            "type": "table",
            "data": read_excel_records(INVOICE_FILE),
        }

    if POLICY_FILE.exists():
        datasets["Finance Policy"] = {
            "file": POLICY_FILE.name,
            "type": "text",
            "data": POLICY_FILE.read_text(encoding="utf-8"),
        }

    return datasets


def get_field(record, *names):
    """Get a field without depending on capitalization."""
    fields = {normalize(key): value for key, value in record.items()}

    for name in names:
        value = fields.get(normalize(name))
        if value is not None:
            return value

    return None


def get_transaction_id(record):
    return get_field(record, "Transaction_id", "Transaction ID")


def get_invoice_id(record):
    return get_field(record, "Invoice_id", "Invoice ID", "Reference")


def get_vendor(record):
    return get_field(record, "Vendor")


def get_amount(record):
    value = get_field(record, "Amount")

    try:
        return float(str(value).replace(",", "").replace("₹", "").strip())
    except (TypeError, ValueError):
        return None


def get_invoice_status(record):
    return normalize(get_field(record, "Status"))


def compare_transaction(transaction, invoice):
    """Return discrepancies between a bank transaction and its invoice."""
    issues = []

    transaction_vendor = normalize(get_vendor(transaction))
    invoice_vendor = normalize(get_vendor(invoice))

    transaction_amount = get_amount(transaction)
    invoice_amount = get_amount(invoice)

    if transaction_vendor != invoice_vendor:
        issues.append(
            f"Vendor mismatch: bank='{get_vendor(transaction)}', "
            f"invoice='{get_vendor(invoice)}'"
        )

    if (
        transaction_amount is not None
        and invoice_amount is not None
        and transaction_amount != invoice_amount
    ):
        difference = transaction_amount - invoice_amount
        issues.append(
            f"Amount mismatch: bank=₹{transaction_amount:,.2f}, "
            f"invoice=₹{invoice_amount:,.2f}, "
            f"difference=₹{difference:,.2f}"
        )

    return issues


def make_match(source, file_name, record, reason=None):
    result = {
        "source": source,
        "file": file_name,
        "type": "record",
        "data": record,
    }

    if reason:
        result["reason"] = reason

    return result


def analyze_transactions(query, bank_records, invoice_records, bank_file, invoice_file):
    """Run structured checks across the complete bank and invoice datasets."""
    q = normalize(query)
    results = []

    invoices_by_id = {
        normalize(get_field(invoice, "Invoice_id", "Invoice ID")): invoice
        for invoice in invoice_records
        if get_field(invoice, "Invoice_id", "Invoice ID") is not None
    }

    is_missing_invoice_search = any(
        phrase in q
        for phrase in (
            "missing invoice",
            "invoice is missing",
            "invoice does not exist",
            "invoice references do not exist",
            "unmatched invoice",
            "no corresponding invoice",
            "invoice reference",
            "orphan transaction",
        )
    )

    is_amount_search = any(
        phrase in q
        for phrase in (
            "amount mismatch",
            "amount difference",
            "amount differs",
            "amounts differ",
            "incorrect amount",
            "discrepanc",
        )
    )

    is_vendor_search = any(
        phrase in q
        for phrase in (
            "vendor mismatch",
            "vendor differs",
            "different vendor",
            "wrong vendor",
        )
    )

    is_general_mismatch = any(
        phrase in q
        for phrase in (
            "mismatch",
            "mismatches",
            "discrepanc",
            "inconsisten",
            "do not match",
            "does not match",
            "don't match",
        )
    )

    is_comparison = any(
        phrase in q
        for phrase in (
            "compare",
            "comparison",
            "matching bank transaction",
            "matches the invoice",
            "match the invoice",
            "check whether",
        )
    )

    is_transaction_search = any(
        phrase in q
        for phrase in (
            "transaction",
            "bank record",
            "bank transaction",
            "pending",
        )
    )

    is_invoice_search = any(
        phrase in q
        for phrase in (
            "invoice",
            "invoices",
        )
    )

    if is_missing_invoice_search:
        for transaction in bank_records:
            reference = normalize(
                get_field(transaction, "Reference", "Invoice_id", "Invoice ID")
            )

            if not reference or reference not in invoices_by_id:
                results.append(
                    make_match(
                        "Bank Transactions",
                        bank_file,
                        transaction,
                        "No matching invoice record was found.",
                    )
                )

        return results

    if is_general_mismatch or is_amount_search or is_vendor_search:
        for transaction in bank_records:
            reference = normalize(
                get_field(transaction, "Reference", "Invoice_id", "Invoice ID")
            )
            invoice = invoices_by_id.get(reference)

            if invoice is None:
                if is_general_mismatch and not is_amount_search and not is_vendor_search:
                    results.append(
                        make_match(
                            "Bank Transactions",
                            bank_file,
                            transaction,
                            "No matching invoice exists.",
                        )
                    )
                continue

            issues = compare_transaction(transaction, invoice)

            if is_amount_search:
                issues = [issue for issue in issues if "Amount mismatch" in issue]

            if is_vendor_search:
                issues = [issue for issue in issues if "Vendor mismatch" in issue]

            if issues:
                results.append(
                    make_match(
                        "Bank Transactions",
                        bank_file,
                        transaction,
                        "; ".join(issues),
                    )
                )

        return results

    # Handle transaction/invoice ID lookups and comparisons.
    ids = re.findall(r"\b(?:INV-\d+|TXN-?\d+)\b", query, flags=re.IGNORECASE)

    if ids:
        wanted_ids = {normalize(item) for item in ids}

        for transaction in bank_records:
            transaction_id = normalize(get_transaction_id(transaction))
            reference = normalize(
                get_field(transaction, "Reference", "Invoice_id", "Invoice ID")
            )

            if transaction_id in wanted_ids or reference in wanted_ids:
                results.append(
                    make_match("Bank Transactions", bank_file, transaction)
                )

        for invoice in invoice_records:
            invoice_id = normalize(get_field(invoice, "Invoice_id", "Invoice ID"))

            if invoice_id in wanted_ids:
                results.append(make_match("Invoices", invoice_file, invoice))

        if is_comparison:
            for transaction in bank_records:
                reference = normalize(
                    get_field(transaction, "Reference", "Invoice_id", "Invoice ID")
                )
                invoice = invoices_by_id.get(reference)

                if (
                    normalize(get_transaction_id(transaction)) in wanted_ids
                    or reference in wanted_ids
                ):
                    if invoice:
                        issues = compare_transaction(transaction, invoice)
                        results.append(
                            make_match(
                                "Comparison",
                                f"{bank_file} + {invoice_file}",
                                {
                                    "Transaction ID": get_transaction_id(transaction),
                                    "Invoice ID": get_field(
                                        invoice, "Invoice_id", "Invoice ID"
                                    ),
                                    "Bank Vendor": get_vendor(transaction),
                                    "Invoice Vendor": get_vendor(invoice),
                                    "Bank Amount": get_amount(transaction),
                                    "Invoice Amount": get_amount(invoice),
                                    "Findings": issues or ["Vendor and amount match."],
                                },
                            )
                        )

        return results

    # Detect vendor names from the actual dataset rather than relying on IDs.
    vendor_query = q
    vendor_found = False

    for record in bank_records + invoice_records:
        vendor = get_vendor(record)

        if vendor and normalize(vendor) in q:
            vendor_query = normalize(vendor)
            vendor_found = True
            break

    if vendor_found:
        if is_invoice_search and not is_transaction_search:
            for invoice in invoice_records:
                if normalize(get_vendor(invoice)) == vendor_query:
                    results.append(make_match("Invoices", invoice_file, invoice))
        elif is_transaction_search and not is_invoice_search:
            for transaction in bank_records:
                if normalize(get_vendor(transaction)) == vendor_query:
                    results.append(
                        make_match("Bank Transactions", bank_file, transaction)
                    )
        else:
            for transaction in bank_records:
                if normalize(get_vendor(transaction)) == vendor_query:
                    results.append(
                        make_match("Bank Transactions", bank_file, transaction)
                    )

            for invoice in invoice_records:
                if normalize(get_vendor(invoice)) == vendor_query:
                    results.append(make_match("Invoices", invoice_file, invoice))

        return results

    # Search general terms against every individual table record.
    stop_words = {
        "find", "show", "search", "look", "up", "all", "the", "for",
        "with", "and", "or", "me", "tell", "report", "list", "identify",
        "which", "what", "are", "is", "about", "belonging", "associated",
        "please", "data", "records", "record", "available", "financepilot",
    }

    terms = [
        term for term in re.findall(r"[a-z0-9₹.]+", q)
        if term not in stop_words and len(term) > 1
    ]

    for dataset_name, records, file_name in (
        ("Bank Transactions", bank_records, bank_file),
        ("Invoices", invoice_records, invoice_file),
    ):
        for record in records:
            searchable_text = normalize(" ".join(
                f"{key} {value}"
                for key, value in record.items()
                if value is not None
            ))

            if terms and all(term in searchable_text for term in terms):
                results.append(make_match(dataset_name, file_name, record))

    return results


def search_files(query):
    """Search data or perform a structured finance analysis."""
    query = str(query).strip()

    if not query:
        return {
            "success": False,
            "query": query,
            "matches": [],
            "total_matches": 0,
            "reason": "Search query is empty.",
        }

    datasets = load_datasets()

    bank_records = datasets.get("Bank Transactions", {}).get("data", [])
    invoice_records = datasets.get("Invoices", {}).get("data", [])

    bank_file = datasets.get("Bank Transactions", {}).get(
        "file", BANK_FILE.name
    )
    invoice_file = datasets.get("Invoices", {}).get(
        "file", INVOICE_FILE.name
    )

    policy_text = datasets.get("Finance Policy", {}).get("data", "")
    q = normalize(query)

    policy_terms = (
        "policy", "approval", "approved", "threshold", "payment rules",
        "manual review", "human approval", "finance rules",
    )

    if policy_text and any(term in q for term in policy_terms):
        matches = [{
            "source": "Finance Policy",
            "file": POLICY_FILE.name,
            "type": "document",
            "data": policy_text,
        }]
    else:
        matches = analyze_transactions(
            query,
            bank_records,
            invoice_records,
            bank_file,
            invoice_file,
        )

    # Fall back to literal searching when structured analysis finds nothing.
    if not matches:
        for dataset_name, dataset in datasets.items():
            if dataset["type"] == "table":
                for record in dataset["data"]:
                    if any(
                        q in normalize(value)
                        for value in record.values()
                        if value is not None
                    ):
                        matches.append(
                            make_match(
                                dataset_name,
                                dataset["file"],
                                record,
                            )
                        )
            elif q in normalize(dataset["data"]):
                matches.append({
                    "source": dataset_name,
                    "file": dataset["file"],
                    "type": "document",
                    "data": dataset["data"],
                })

    return {
        "success": True,
        "query": query,
        "matches": matches,
        "total_matches": len(matches),
    }


def format_search_results(search_result):
    """Format search results for display."""
    if not search_result.get("success"):
        return f"Search failed: {search_result.get('reason', 'Unknown error')}"

    lines = [
        f"Search Query: {search_result.get('query', '')}",
        f"Matches Found: {search_result.get('total_matches', 0)}",
        "",
    ]

    matches = search_result.get("matches", [])

    if not matches:
        lines.append(
            "No matching records were found. Try a vendor name, "
            "invoice reference, transaction ID, or a more specific question."
        )
        return "\n".join(lines)

    for number, match in enumerate(matches, start=1):
        lines.extend([
            f"Match {number}",
            f"Source: {match.get('source', 'Unknown')}",
            f"File: {match.get('file', 'Unknown')}",
        ])

        if match.get("reason"):
            lines.append(f"Finding: {match['reason']}")

        data = match.get("data", {})

        if isinstance(data, dict):
            for key, value in data.items():
                lines.append(f"{key}: {value}")
        else:
            lines.append(str(data))

        lines.append("")

    return "\n".join(lines)


if __name__ == "__main__":
    print("FinancePilot File Search Test")
    print("=============================")

    while True:
        query = input("\nEnter a query (or type exit): ").strip()

        if query.casefold() == "exit":
            break

        result = search_files(query)
        print()
        print(format_search_results(result))
